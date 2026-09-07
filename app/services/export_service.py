"""Excel export builders for journals, weekly summaries, and monthly rollups."""

from __future__ import annotations

from calendar import monthrange
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.core.exceptions import BadRequestError
from app.models.journal_entry import JournalEntry
from app.models.weekly_summary import WeeklySummary
from app.services import analytics_service

_CONTENT_MAX = 32000
_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _workbook_bytes(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _autosize(ws, max_width: int = 48) -> None:
    for idx, column_cells in enumerate(ws.columns, start=1):
        length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            length = max(length, min(len(value), max_width))
        ws.column_dimensions[get_column_letter(idx)].width = max(length + 2, 12)


def _header_row(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)


def _truncate(text: str | None) -> str:
    if not text:
        return ""
    if len(text) <= _CONTENT_MAX:
        return text
    return f"{text[: _CONTENT_MAX - 1]}…"


def _suggestions_text(suggestions) -> str:
    if not suggestions:
        return ""
    if isinstance(suggestions, list):
        return " | ".join(str(item) for item in suggestions)
    return str(suggestions)


def _validate_range(date_from: date | None, date_to: date | None) -> tuple[date | None, date | None]:
    if date_from and date_to and date_from > date_to:
        raise BadRequestError("date_from must be on or before date_to")
    return date_from, date_to


def _day_start(day: date) -> datetime:
    return datetime.combine(day, time.min)


def _day_end(day: date) -> datetime:
    return datetime.combine(day, time.max)


def _load_journals(
    db: Session,
    user_id: int,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> list[JournalEntry]:
    stmt = (
        select(JournalEntry)
        .options(joinedload(JournalEntry.sentiment))
        .where(JournalEntry.user_id == user_id)
        .order_by(JournalEntry.created_at.asc())
    )
    if date_from is not None:
        stmt = stmt.where(JournalEntry.created_at >= _day_start(date_from))
    if date_to is not None:
        stmt = stmt.where(JournalEntry.created_at <= _day_end(date_to))
    return list(db.scalars(stmt).unique().all())


def _write_journals_sheet(ws, journals: list[JournalEntry]) -> None:
    _header_row(
        ws,
        [
            "ID",
            "Title",
            "Content",
            "Created At",
            "Sentiment",
            "Mood",
            "Emotion",
            "Confidence",
        ],
    )
    for journal in journals:
        sentiment = journal.sentiment
        ws.append(
            [
                journal.id,
                journal.title or "",
                _truncate(journal.content),
                journal.created_at.isoformat(sep=" ", timespec="seconds") if journal.created_at else "",
                sentiment.sentiment if sentiment else "",
                sentiment.mood if sentiment else "",
                sentiment.emotion if sentiment else "",
                sentiment.confidence if sentiment else "",
            ]
        )
    _autosize(ws)


def _load_weekly_summaries(
    db: Session,
    user_id: int,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> list[WeeklySummary]:
    stmt = (
        select(WeeklySummary)
        .where(WeeklySummary.user_id == user_id)
        .order_by(WeeklySummary.week_start.asc())
    )
    if date_from is not None:
        stmt = stmt.where(WeeklySummary.week_end >= date_from)
    if date_to is not None:
        stmt = stmt.where(WeeklySummary.week_start <= date_to)
    return list(db.scalars(stmt).all())


def _write_summaries_sheet(ws, summaries: list[WeeklySummary]) -> None:
    _header_row(
        ws,
        [
            "ID",
            "Week Start",
            "Week End",
            "Entry Count",
            "Summary",
            "Suggestions",
        ],
    )
    for summary in summaries:
        ws.append(
            [
                summary.id,
                summary.week_start.isoformat(),
                summary.week_end.isoformat(),
                summary.entry_count,
                _truncate(summary.summary),
                _suggestions_text(summary.suggestions),
            ]
        )
    _autosize(ws)


def export_journals(
    db: Session,
    user_id: int,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> tuple[bytes, str]:
    date_from, date_to = _validate_range(date_from, date_to)
    journals = _load_journals(db, user_id, date_from=date_from, date_to=date_to)

    wb = Workbook()
    ws = wb.active
    ws.title = "Journals"
    _write_journals_sheet(ws, journals)

    from_label = date_from.isoformat() if date_from else "all"
    to_label = date_to.isoformat() if date_to else "all"
    filename = f"sentitrack-journals-{from_label}-{to_label}.xlsx"
    return _workbook_bytes(wb), filename


def export_weekly_summaries(
    db: Session,
    user_id: int,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> tuple[bytes, str]:
    date_from, date_to = _validate_range(date_from, date_to)
    summaries = _load_weekly_summaries(db, user_id, date_from=date_from, date_to=date_to)

    wb = Workbook()
    ws = wb.active
    ws.title = "Weekly Summaries"
    _write_summaries_sheet(ws, summaries)

    from_label = date_from.isoformat() if date_from else "all"
    to_label = date_to.isoformat() if date_to else "all"
    filename = f"sentitrack-weekly-summaries-{from_label}-{to_label}.xlsx"
    return _workbook_bytes(wb), filename


def export_monthly_summary(
    db: Session,
    user_id: int,
    *,
    year: int,
    month: int,
) -> tuple[bytes, str]:
    if month < 1 or month > 12:
        raise BadRequestError("month must be between 1 and 12")

    start = date(year, month, 1)
    end = date(year, month, monthrange(year, month)[1])

    journals = _load_journals(db, user_id, date_from=start, date_to=end)
    summaries = _load_weekly_summaries(db, user_id, date_from=start, date_to=end)

    # KPIs via the same analytics helpers used by the dashboard.
    rows = [
        analytics_service._Row(
            j.created_at,
            j.sentiment.sentiment if j.sentiment else None,
            j.sentiment.emotion if j.sentiment else None,
            j.sentiment.mood if j.sentiment else None,
            j.sentiment.confidence if j.sentiment else None,
        )
        for j in journals
    ]
    sentiment_counts = analytics_service._sentiment_counts(rows)
    analyzed = sum(1 for r in rows if r.sentiment)

    wb = Workbook()

    kpi = wb.active
    kpi.title = "Month KPIs"
    _header_row(kpi, ["Metric", "Value"])
    kpi_rows = [
        ("Year", year),
        ("Month", month),
        ("Range Start", start.isoformat()),
        ("Range End", end.isoformat()),
        ("Entries", len(journals)),
        ("Analyzed", analyzed),
        ("Positive", sentiment_counts.get("positive", 0)),
        ("Neutral", sentiment_counts.get("neutral", 0)),
        ("Negative", sentiment_counts.get("negative", 0)),
        ("Average Confidence", analytics_service._average_confidence(rows) or ""),
        ("Top Mood", analytics_service._most_common(r.mood for r in rows) or ""),
        ("Top Emotion", analytics_service._most_common(r.emotion for r in rows) or ""),
        ("Weekly Summaries", len(summaries)),
    ]
    for metric, value in kpi_rows:
        kpi.append([metric, value])
    _autosize(kpi)

    journals_sheet = wb.create_sheet("Journals")
    _write_journals_sheet(journals_sheet, journals)

    summaries_sheet = wb.create_sheet("Weekly Summaries")
    _write_summaries_sheet(summaries_sheet, summaries)

    filename = f"sentitrack-monthly-summary-{year:04d}-{month:02d}.xlsx"
    return _workbook_bytes(wb), filename


# Re-export media type for routes.
XLSX_MEDIA_TYPE = _XLSX_MEDIA_TYPE
