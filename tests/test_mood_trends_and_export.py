"""Smoke tests for mood-trends aggregation and Excel export endpoints."""

from __future__ import annotations

from datetime import date, datetime, timedelta
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.security import create_access_token, hash_password
from app.database import Base, get_db
from app.main import create_app
from app.models.journal_entry import JournalEntry
from app.models.sentiment import Sentiment
from app.models.user import User
from app.models.weekly_summary import WeeklySummary
from app.services import export_service
from app.services.analytics_service import _Row, _build_mood_trends


def test_build_mood_trends_fills_empty_days_and_top_emotions():
    monday = date(2026, 9, 7)  # Monday
    rows = [
        _Row(datetime(2026, 9, 7, 10, 0), "positive", "Happy", "upbeat", 0.9),
        _Row(datetime(2026, 9, 7, 18, 0), "negative", "sad", "low", 0.8),
        _Row(datetime(2026, 9, 9, 12, 0), "negative", "Sad", "down", 0.7),
        _Row(datetime(2026, 9, 9, 15, 0), "positive", "angry", "tense", 0.6),
        _Row(datetime(2026, 9, 10, 9, 0), None, None, None, None),
    ]

    result = _build_mood_trends(
        rows,
        period="week",
        start=monday,
        end=monday + timedelta(days=6),
        top_emotions=2,
    )

    assert result["period"] == "week"
    assert len(result["buckets"]) == 7
    assert result["series"]["emotions"] == ["sad", "happy"]

    monday_bucket = result["buckets"][0]
    assert monday_bucket["date"] == monday
    assert monday_bucket["sentiment"] == {"positive": 1, "neutral": 0, "negative": 1}
    assert monday_bucket["emotions"]["happy"] == 1
    assert monday_bucket["emotions"]["sad"] == 1

    empty_tuesday = result["buckets"][1]
    assert empty_tuesday["sentiment"] == {"positive": 0, "neutral": 0, "negative": 0}
    assert empty_tuesday["emotions"] == {"sad": 0, "happy": 0}

    assert result["totals"]["analyzed"] == 4
    assert result["totals"]["entries"] == 5


@pytest.fixture()
def client_and_user():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    Base.metadata.create_all(bind=engine)

    db = TestingSessionLocal()
    user = User(
        name="Smoke Tester",
        email="smoke@example.com",
        password_hash=hash_password("password123"),
    )
    db.add(user)
    db.commit()
    db.refresh(user)

    today = date.today()
    week_start = today - timedelta(days=today.weekday())
    journal = JournalEntry(
        user_id=user.id,
        title="Day one",
        content="Feeling okay overall.",
        created_at=datetime.combine(week_start, datetime.min.time()),
    )
    db.add(journal)
    db.flush()
    db.add(
        Sentiment(
            journal_id=journal.id,
            sentiment="positive",
            mood="calm",
            emotion="happy",
            confidence=0.91,
            raw_response="{}",
        )
    )
    db.add(
        WeeklySummary(
            user_id=user.id,
            week_start=week_start,
            week_end=week_start + timedelta(days=6),
            summary="A steady week.",
            suggestions=["Keep journaling"],
            entry_count=1,
        )
    )
    db.commit()

    app = create_app()

    def override_get_db():
        try:
            yield db
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)
    token = create_access_token(user.id)

    try:
        yield client, token, user, week_start
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(bind=engine)
        engine.dispose()


def test_export_requires_auth(client_and_user):
    client, _token, _user, _week_start = client_and_user
    response = client.get("/api/v1/export/journals")
    assert response.status_code == 401


def test_export_journals_returns_xlsx(client_and_user):
    client, token, _user, week_start = client_and_user
    response = client.get(
        "/api/v1/export/journals",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment;" in response.headers["content-disposition"]
    assert "sentitrack-journals-" in response.headers["content-disposition"]
    assert len(response.content) > 100

    wb = load_workbook(BytesIO(response.content))
    assert wb.active.title == "Journals"
    assert wb.active["A2"].value == 1


def test_export_monthly_summary_workbook(client_and_user):
    client, token, _user, week_start = client_and_user
    response = client.get(
        f"/api/v1/export/monthly-summary?year={week_start.year}&month={week_start.month}",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = load_workbook(BytesIO(response.content))
    assert set(wb.sheetnames) == {"Month KPIs", "Journals", "Weekly Summaries"}


def test_mood_trends_endpoint(client_and_user):
    client, token, _user, week_start = client_and_user
    response = client.get(
        f"/api/v1/analytics/mood-trends?period=week&anchor={week_start.isoformat()}",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["period"] == "week"
    assert len(body["buckets"]) == 7
    assert body["series"]["emotions"] == ["happy"]
    assert body["totals"]["analyzed"] == 1


def test_export_service_bytes_nonempty():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Session = sessionmaker(bind=engine)
    Base.metadata.create_all(bind=engine)
    db = Session()
    try:
        payload, filename = export_service.export_weekly_summaries(db, user_id=99999)
        assert filename.startswith("sentitrack-weekly-summaries-")
        assert len(payload) > 50
        wb = load_workbook(BytesIO(payload))
        assert wb.active.title == "Weekly Summaries"
    finally:
        db.close()
        Base.metadata.drop_all(bind=engine)
        engine.dispose()
